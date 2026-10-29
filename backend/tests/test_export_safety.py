"""
tests/test_export_safety.py — an exported file must not carry live
formulas from the uploaded data into the recipient's spreadsheet.
"""
from __future__ import annotations

import io

import openpyxl
import pandas as pd
import pytest

PAYLOADS = ['=HYPERLINK("http://evil.example","click")', "+1+1", "-2+3", "@SUM(A1)"]


@pytest.fixture()
def risky_dataset(client):
    df = pd.DataFrame({"note": PAYLOADS + ["plain", "-5"], "v": range(6)})
    r = client.post("/api/datasets/upload", files={
        "file": ("x.csv", io.BytesIO(df.to_csv(index=False).encode()), "text/csv")})
    ds = r.json()["meta"]["dataset_id"]
    yield ds
    client.delete(f"/api/datasets/{ds}")


def test_excel_export_holds_no_formulas(client, risky_dataset):
    wb = openpyxl.load_workbook(io.BytesIO(client.get(f"/api/reports/{risky_dataset}/excel").content))
    cells = [row[0] for row in wb["Data"].iter_rows(min_row=2, max_col=1)]
    assert all(c.data_type == "s" for c in cells)
    assert [c.value for c in cells][:4] == PAYLOADS      # shown exactly as uploaded


def test_csv_export_neutralises_formula_cells_but_not_numbers(client, risky_dataset):
    text = client.get(f"/api/reports/{risky_dataset}/csv").content.decode("utf-8-sig")
    notes = pd.read_csv(io.StringIO(text), dtype=str)["note"].tolist()
    assert notes[:4] == ["'" + p for p in PAYLOADS]
    assert notes[4:] == ["plain", "-5"]


def test_numeric_columns_are_exported_unchanged(client, risky_dataset):
    text = client.get(f"/api/reports/{risky_dataset}/csv").content.decode("utf-8-sig")
    assert pd.read_csv(io.StringIO(text))["v"].tolist() == list(range(6))


def test_excel_export_handles_timezone_aware_timestamps(client):
    df = pd.DataFrame({"logged": pd.date_range("2024-01-01", periods=5, freq="h", tz="UTC"),
                       "v": range(5)})
    r = client.post("/api/datasets/upload", files={
        "file": ("tz.csv", io.BytesIO(df.to_csv(index=False).encode()), "text/csv")})
    ds = r.json()["meta"]["dataset_id"]
    try:
        x = client.get(f"/api/reports/{ds}/excel")
        assert x.status_code == 200
        ws = openpyxl.load_workbook(io.BytesIO(x.content))["Data"]
        assert ws.cell(1, 1).value == "logged (UTC)"
        assert str(ws.cell(2, 1).value).startswith("2024-01-01 00:00")
    finally:
        client.delete(f"/api/datasets/{ds}")
